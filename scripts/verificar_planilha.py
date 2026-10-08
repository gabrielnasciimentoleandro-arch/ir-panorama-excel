#!/usr/bin/env python3
"""Validação estrutural, funcional e de privacidade do IR Panorama."""

from __future__ import annotations

from math import isclose
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
FILE = ROOT / "ir_panorama_organizador.xlsx"
AUTHOR = "0Barone (@gabrielnasciimentoleandro-arch)"

EXPECTED_SHEETS = [
    "Início", "Titular", "Dependentes", "Rendimentos", "Bens",
    "Pagamentos", "Informes", "Notas", "Guia", "Apoio",
]
EXPECTED_NAMES = {
    "ano_calendario", "exercicio_declaracao", "nome_titular", "cpf_titular",
    "dados_rendimentos", "dados_bens", "dados_pagamentos", "dados_informes",
    "lista_status_documento", "lista_categorias_rendimento", "lista_grupos_bens",
    "lista_categorias_pagamento", "lista_prioridades", "lista_status_nota",
    "lista_estados_civis", "lista_ufs", "lista_relacoes_dependente",
    "lista_sim_nao", "lista_titular_dependente", "lista_categorias_informe",
    "lista_tipos_declaracao", "lista_status_cadastro", "lista_categorias_nota",
    "lista_paises",
}


def brl(value: float) -> str:
    text = f"{value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"R$ {text}"


def nonempty_ids(ws) -> int:
    return sum(1 for row in range(9, 39) if ws.cell(row, 1).value)


def main() -> None:
    assert FILE.exists(), f"Arquivo não encontrado: {FILE}"

    with ZipFile(FILE) as zf:
        assert zf.testzip() is None, "Há uma entrada corrompida no XLSX"
        names = set(zf.namelist())
        assert "xl/workbook.xml" in names
        assert "docProps/core.xml" in names
        assert any(name.startswith("xl/charts/chart") for name in names), "Gráfico ausente"
        assert not any(name.startswith("xl/externalLinks/") for name in names), "Vínculo externo encontrado"
        assert not any(name.endswith("vbaProject.bin") for name in names), "Macro encontrada"

    wb = load_workbook(FILE, data_only=False)
    cached = load_workbook(FILE, data_only=True)

    assert wb.sheetnames == EXPECTED_SHEETS
    assert wb["Apoio"].sheet_state == "hidden"
    assert wb.properties.creator == AUTHOR
    assert EXPECTED_NAMES == set(wb.defined_names), "Intervalos nomeados divergentes"

    table_count = sum(len(ws.tables) for ws in wb.worksheets)
    validation_count = sum(len(ws.data_validations.dataValidation) for ws in wb.worksheets)
    chart_count = sum(len(ws._charts) for ws in wb.worksheets)
    assert table_count == 6, f"Esperadas 6 tabelas; encontradas {table_count}"
    assert validation_count >= 25, f"Poucas validações: {validation_count}"
    assert chart_count >= 1

    inicio = wb["Início"]
    assert inicio["A10"].value == "=SUM(Rendimentos!$E$9:$E$38)"
    assert inicio["D10"].value == "=SUM(Bens!$H$9:$H$38)"
    assert inicio["G10"].value == "=SUM(Pagamentos!$H$9:$H$38)"
    assert "COUNTIF(Informes!$H$9:$H$38" in inicio["J10"].value
    assert wb["Pagamentos"]["H9"].value == '=IF(COUNTA(F9:G9)=0,"",MAX(0,F9-G9))'
    assert "TODAY()" in wb["Notas"]["H9"].value

    internal_links = 0
    official_links = set()
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if not cell.hyperlink:
                    continue
                target = cell.hyperlink.target
                if target.startswith("#"):
                    internal_links += 1
                elif target.startswith("https://"):
                    official_links.add(target)
    assert internal_links >= 80, f"Poucos links internos: {internal_links}"
    assert len(official_links) == 3, "Links oficiais inesperados"
    assert all("gov.br" in url or "fazenda.gov.br" in url for url in official_links)

    # Resultados do exemplo, lidos do cache OOXML.
    cstart = cached["Início"]
    assert isclose(cstart["A10"].value, 93830.75, abs_tol=0.001)
    assert isclose(cstart["D10"].value, 337650.75, abs_tol=0.001)
    assert isclose(cstart["G10"].value, 12460.00, abs_tol=0.001)
    assert isclose(cstart["J10"].value, 0.625, abs_tol=1e-9)
    assert isclose(cstart["A15"].value, 7800.00, abs_tol=0.001)
    assert isclose(cstart["D15"].value, 7260.00, abs_tol=0.001)
    assert cstart["G15"].value == 5
    assert cstart["J15"].value == 11
    assert cstart["O3"].value == 5
    assert cstart["O4"].value == 2
    assert cstart["O5"].value == 1
    assert cstart["O6"].value == 0

    assert nonempty_ids(wb["Dependentes"]) == 1
    assert nonempty_ids(wb["Rendimentos"]) == 4
    assert nonempty_ids(wb["Bens"]) == 4
    assert nonempty_ids(wb["Pagamentos"]) == 3
    assert nonempty_ids(wb["Informes"]) == 8
    assert nonempty_ids(wb["Notas"]) == 3

    allowed_doc_status = {"Pendente", "Recebido", "Conferido", "Não se aplica"}
    for sheet, col in [("Rendimentos", 9), ("Bens", 10), ("Pagamentos", 10), ("Informes", 8)]:
        ws = wb[sheet]
        for row in range(9, 39):
            if ws.cell(row, 1).value:
                assert ws.cell(row, col).value in allowed_doc_status

    # Identificadores de demonstração não podem parecer dados pessoais reais.
    assert wb["Titular"]["C9"].value == "000.000.000-00"
    assert wb["Titular"]["C11"].value.endswith(".invalid")
    for sheet, column in [("Dependentes", 3), ("Rendimentos", 4), ("Pagamentos", 4), ("Informes", 4)]:
        ws = wb[sheet]
        for row in range(9, 39):
            value = ws.cell(row, column).value
            if value and value not in {"N/A"}:
                assert "000.000" in str(value), f"Identificador não fictício em {sheet}!{ws.cell(row, column).coordinate}"

    # Não deve existir erro armazenado no exemplo calculado.
    excel_errors = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!"}
    for ws in cached.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                assert cell.value not in excel_errors, f"Erro em {ws.title}!{cell.coordinate}: {cell.value}"

    for name in ("painel-geral.png", "controle-informes.png"):
        path = ROOT / "docs" / name
        assert path.exists(), f"Prévia ausente: {name}"
        with Image.open(path) as image:
            assert image.width >= 1500 and image.height >= 850

    print("VALIDAÇÃO APROVADA")
    print(f"Arquivo: {FILE.name}")
    print(f"Autoria: {wb.properties.creator}")
    print(f"Abas: {len(wb.sheetnames)} ({', '.join(wb.sheetnames)})")
    print(f"Tabelas: {table_count} | Validações de dados: {validation_count} | Gráficos: {chart_count}")
    print(f"Intervalos nomeados: {len(wb.defined_names)}")
    print(f"Links internos: {internal_links} | Links oficiais: {len(official_links)}")
    print(f"Rendimentos do exemplo: {brl(cstart['A10'].value)}")
    print(f"Bens na situação atual: {brl(cstart['D10'].value)}")
    print(f"Pagamentos líquidos: {brl(cstart['G10'].value)}")
    print(f"Informes conferidos: {cstart['J10'].value:.1%}")
    print("Integridade ZIP/XML, fórmulas, caches, tabelas, listas e privacidade: OK")
    print("Macros e vínculos externos de pasta de trabalho: AUSENTES")


if __name__ == "__main__":
    main()
