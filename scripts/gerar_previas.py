#!/usr/bin/env python3
"""Gera imagens de demonstração do IR Panorama a partir da planilha pronta."""

from __future__ import annotations

from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, Rectangle
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
FILE = ROOT / "ir_panorama_organizador.xlsx"
DOCS = ROOT / "docs"
AUTHOR = "0Barone (@gabrielnasciimentoleandro-arch)"

NAVY = "#0F172A"
NAVY_2 = "#1E293B"
SLATE = "#334155"
CYAN = "#38BDF8"
CYAN_DARK = "#0284C7"
TEAL = "#14B8A6"
EMERALD = "#10B981"
AMBER = "#F59E0B"
RED = "#EF4444"
WHITE = "#FFFFFF"
TEXT = "#172033"
MUTED = "#64748B"
LIGHT = "#F8FAFC"
BORDER = "#CBD5E1"
PALE_GREEN = "#DCFCE7"
PALE_BLUE = "#ECFEFF"
PALE_AMBER = "#FEF3C7"
PALE_RED = "#FEE2E2"


def brl(value: float) -> str:
    text = f"{value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"R$ {text}"


def pct(value: float) -> str:
    return f"{value * 100:.0f}%"


def box(ax, x, y, w, h, face=WHITE, edge=BORDER, radius=.014, lw=1):
    patch = FancyBboxPatch(
        (x, y), w, h,
        boxstyle=f"round,pad=0.004,rounding_size={radius}",
        facecolor=face, edgecolor=edge, linewidth=lw,
        transform=ax.transAxes, clip_on=False,
    )
    ax.add_patch(patch)
    return patch


def dashboard_preview(wb) -> None:
    ws = wb["Início"]
    income = ws["A10"].value
    assets = ws["D10"].value
    payments = ws["G10"].value
    progress = ws["J10"].value
    irrf = ws["A15"].value
    inss = ws["D15"].value
    pending = ws["G15"].value
    records = ws["J15"].value
    statuses = [ws[f"O{row}"].value for row in range(3, 7)]

    fig = plt.figure(figsize=(16, 10), facecolor="#E8EEF5")
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_axis_off()
    box(ax, .02, .02, .96, .96, WHITE, "#BFCADA", .012)

    ax.add_patch(Rectangle((.02, .865), .96, .115, transform=ax.transAxes,
                           facecolor=NAVY, edgecolor=NAVY))
    ax.text(.05, .925, "IR PANORAMA", transform=ax.transAxes, color=WHITE,
            fontsize=27, fontweight="bold", va="center")
    ax.text(.05, .885, "Organizador de documentos, dados e pendências do Imposto de Renda",
            transform=ax.transAxes, color="#BED0E5", fontsize=11, va="center")
    ax.text(.94, .925, "DADOS FICTÍCIOS", transform=ax.transAxes, color=NAVY,
            fontsize=9, fontweight="bold", ha="right", va="center",
            bbox=dict(boxstyle="round,pad=.5", facecolor=CYAN, edgecolor=CYAN))
    ax.text(.94, .885, "Exercício 2026", transform=ax.transAxes, color=WHITE,
            fontsize=10, ha="right", va="center")

    cards = [
        (.05, "RENDIMENTOS BRUTOS", brl(income), CYAN_DARK),
        (.285, "BENS — SITUAÇÃO ATUAL", brl(assets), EMERALD),
        (.52, "PAGAMENTOS LÍQUIDOS", brl(payments), AMBER),
        (.755, "DOCUMENTOS CONFERIDOS", pct(progress), TEAL),
    ]
    for x, label, value, accent in cards:
        box(ax, x, .685, .195, .14, WHITE)
        ax.add_patch(Rectangle((x, .785), .195, .04, transform=ax.transAxes,
                               facecolor=NAVY_2, edgecolor=NAVY_2))
        ax.text(x + .0975, .805, label, transform=ax.transAxes, color=WHITE,
                fontsize=8.3, fontweight="bold", ha="center", va="center")
        ax.text(x + .0975, .735, value, transform=ax.transAxes, color=accent,
                fontsize=18, fontweight="bold", ha="center", va="center")

    small_cards = [
        ("IR retido", brl(irrf), CYAN_DARK),
        ("INSS informado", brl(inss), EMERALD),
        ("Pendências abertas", str(pending), RED),
        ("Registros financeiros", str(records), SLATE),
    ]
    sx = .05
    for label, value, accent in small_cards:
        box(ax, sx, .59, .195, .065, LIGHT)
        ax.text(sx + .015, .625, label, transform=ax.transAxes, color=MUTED,
                fontsize=8.5, va="center")
        ax.text(sx + .18, .625, value, transform=ax.transAxes, color=accent,
                fontsize=11, fontweight="bold", ha="right", va="center")
        sx += .235

    ax.text(.05, .552, "PROGRESSO DA CONFERÊNCIA DOCUMENTAL", transform=ax.transAxes,
            fontsize=9, fontweight="bold", color=TEXT)
    box(ax, .05, .515, .9, .025, "#E2E8F0", "#E2E8F0", .012)
    ax.add_patch(FancyBboxPatch(
        (.05, .515), .9 * progress, .025,
        boxstyle="round,pad=0.004,rounding_size=.012",
        facecolor=TEAL, edgecolor=TEAL, transform=ax.transAxes,
    ))
    ax.text(.955, .528, pct(progress), transform=ax.transAxes, color=TEXT,
            fontsize=10, fontweight="bold", ha="right", va="center")

    ax.text(.05, .475, "MENU PRINCIPAL", transform=ax.transAxes, color=TEXT,
            fontsize=11, fontweight="bold")
    menus = [
        ("01", "TITULAR", CYAN), ("02", "DEPENDENTES", TEAL),
        ("03", "RENDIMENTOS", EMERALD), ("04", "BENS", AMBER),
        ("05", "PAGAMENTOS", AMBER), ("06", "INFORMES", TEAL),
        ("07", "NOTAS", RED), ("08", "GUIA", CYAN),
    ]
    for index, (number, label, accent) in enumerate(menus):
        row = index // 4
        col = index % 4
        x = .05 + col * .225
        y = .375 - row * .105
        box(ax, x, y, .205, .075, WHITE)
        ax.add_patch(Rectangle((x, y), .008, .075, transform=ax.transAxes,
                               facecolor=accent, edgecolor=accent))
        ax.text(x + .02, y + .047, number, transform=ax.transAxes, color=accent,
                fontsize=9, fontweight="bold", va="center")
        ax.text(x + .06, y + .047, label, transform=ax.transAxes, color=TEXT,
                fontsize=9, fontweight="bold", va="center")
        ax.text(x + .02, y + .018, "Abrir área", transform=ax.transAxes, color=MUTED,
                fontsize=7.5, va="center")

    ax.text(.05, .185, "RESUMO DOS INFORMES", transform=ax.transAxes, color=TEXT,
            fontsize=11, fontweight="bold")
    pie_ax = fig.add_axes([.055, .045, .18, .145])
    pie_ax.pie(statuses, colors=[EMERALD, CYAN, RED, "#94A3B8"], startangle=90,
               wedgeprops=dict(width=.35, edgecolor=WHITE))
    pie_ax.text(0, .06, f"{sum(statuses)}", ha="center", va="center",
                color=TEXT, fontsize=16, fontweight="bold")
    pie_ax.text(0, -.2, "informes", ha="center", va="center", color=MUTED, fontsize=8)
    pie_ax.set_axis_off()

    legend = [
        ("Conferidos", statuses[0], EMERALD),
        ("Recebidos", statuses[1], CYAN),
        ("Pendentes", statuses[2], RED),
        ("Não se aplica", statuses[3], "#94A3B8"),
    ]
    lx, ly = .25, .16
    for i, (label, value, color) in enumerate(legend):
        x = lx + (i % 2) * .22
        y = ly - (i // 2) * .06
        ax.add_patch(Rectangle((x, y - .012), .016, .016, transform=ax.transAxes,
                               facecolor=color, edgecolor=color))
        ax.text(x + .025, y, f"{label}: {value}", transform=ax.transAxes,
                color=TEXT, fontsize=9, va="center")

    box(ax, .7, .06, .25, .13, PALE_AMBER, "#F5CC77")
    ax.text(.72, .165, "PRIVACIDADE", transform=ax.transAxes, color=TEXT,
            fontsize=10, fontweight="bold")
    ax.text(.72, .135,
            "Publique somente o modelo\ncom dados fictícios. Mantenha sua\ncópia real em local privado e protegido.",
            transform=ax.transAxes, color=TEXT, fontsize=8.5, va="top")

    ax.text(.05, .035, f"Projeto autoral • {AUTHOR} • sem macros, APIs ou conexões externas",
            transform=ax.transAxes, color=MUTED, fontsize=8.5)

    output = DOCS / "painel-geral.png"
    fig.savefig(output, dpi=160, bbox_inches="tight", facecolor=fig.get_facecolor())
    plt.close(fig)
    print(f"Prévia criada: {output}")


def informes_preview(wb) -> None:
    ws = wb["Informes"]
    rows = []
    for row in range(9, 39):
        if not ws.cell(row, 1).value:
            continue
        rows.append([ws.cell(row, col).value for col in range(1, 10)])

    status_colors = {
        "Conferido": (PALE_GREEN, EMERALD),
        "Recebido": (PALE_BLUE, CYAN_DARK),
        "Pendente": (PALE_RED, RED),
        "Não se aplica": ("#E2E8F0", MUTED),
    }

    fig = plt.figure(figsize=(16, 9), facecolor="#E8EEF5")
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_axis_off()
    box(ax, .02, .025, .96, .95, WHITE, "#BFCADA", .012)
    ax.add_patch(Rectangle((.02, .84), .96, .135, transform=ax.transAxes,
                           facecolor=NAVY, edgecolor=NAVY))
    ax.text(.05, .915, "IR PANORAMA  /  INFORMES", transform=ax.transAxes,
            color=WHITE, fontsize=24, fontweight="bold", va="center")
    ax.text(.05, .867, "Checklist central de documentos — demonstração com dados fictícios",
            transform=ax.transAxes, color="#BED0E5", fontsize=11, va="center")

    completed = sum(1 for row in rows if row[7] in ("Conferido", "Não se aplica"))
    progress = completed / len(rows)
    ax.text(.94, .915, pct(progress), transform=ax.transAxes, color=CYAN,
            fontsize=22, fontweight="bold", ha="right", va="center")
    ax.text(.94, .875, "conferência concluída", transform=ax.transAxes, color=WHITE,
            fontsize=9, ha="right", va="center")

    headers = ["ID", "Categoria", "Fonte/Instituição", "Documento esperado", "Recebido em", "Status"]
    col_x = [.055, .135, .265, .53, .75, .855]
    col_w = [.07, .12, .255, .21, .095, .09]
    header_y = .785
    for x, w, label in zip(col_x, col_w, headers):
        ax.add_patch(Rectangle((x, header_y), w, .045, transform=ax.transAxes,
                               facecolor=NAVY_2, edgecolor=WHITE, linewidth=.5))
        ax.text(x + .006, header_y + .0225, label, transform=ax.transAxes,
                color=WHITE, fontsize=8.5, fontweight="bold", va="center")

    y = .72
    row_h = .077
    for idx, row in enumerate(rows):
        face = WHITE if idx % 2 == 0 else LIGHT
        ax.add_patch(Rectangle((.055, y), .89, row_h - .006, transform=ax.transAxes,
                               facecolor=face, edgecolor=BORDER, linewidth=.6))
        received = row[6].strftime("%d/%m/%Y") if row[6] else "—"
        values = [row[0], row[1], row[2], row[4], received]
        for x, value in zip(col_x[:5], values):
            text = str(value)
            if len(text) > 32:
                text = text[:30] + "…"
            ax.text(x + .006, y + (row_h - .006) / 2, text, transform=ax.transAxes,
                    color=TEXT, fontsize=8.3, va="center")
        status = row[7]
        bg, color = status_colors[status]
        ax.text(col_x[5] + col_w[5] / 2, y + (row_h - .006) / 2, status,
                transform=ax.transAxes, color=color, fontsize=8.3, fontweight="bold",
                ha="center", va="center",
                bbox=dict(boxstyle="round,pad=.35", facecolor=bg, edgecolor=bg))
        y -= row_h

    ax.text(.055, .075, "Fluxo sugerido: Pendente  →  Recebido  →  Conferido",
            transform=ax.transAxes, color=TEXT, fontsize=10, fontweight="bold")
    ax.text(.055, .04,
            f"Projeto autoral • {AUTHOR} • nenhum documento real está incluído nesta demonstração",
            transform=ax.transAxes, color=MUTED, fontsize=8.5)

    output = DOCS / "controle-informes.png"
    fig.savefig(output, dpi=160, bbox_inches="tight", facecolor=fig.get_facecolor())
    plt.close(fig)
    print(f"Prévia criada: {output}")


def main() -> None:
    DOCS.mkdir(parents=True, exist_ok=True)
    if not FILE.exists():
        raise SystemExit("Gere a planilha antes de criar as prévias.")
    wb = load_workbook(FILE, data_only=True)
    dashboard_preview(wb)
    informes_preview(wb)


if __name__ == "__main__":
    main()
